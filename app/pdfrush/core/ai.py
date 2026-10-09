"""PDF Intelligence: AI Summarizer, Translate PDF and Ask-your-PDF.

Uses the online AI (free Gemini/Groq keys, Claude, OpenAI...) or the offline local model.
The summarizer still works with no AI at all (extractive summary)."""
from __future__ import annotations

import json
import re
import tempfile
from collections import Counter
from pathlib import Path

import pypdfium2 as pdfium

from .. import paths
from . import llm, textlayer
from .registry import PDF, Param, tool
from .util import (hex_to_rgb, open_pdf, out_name, overlay_pages, parse_pages, rect_to_visual, save, script_of,
                   unicode_font)

INTEL = "PDF Intelligence"
LANGUAGES = ["English", "French", "Spanish", "Portuguese", "German", "Italian", "Dutch", "Swahili", "Luganda",
             "Kinyarwanda", "Amharic", "Somali", "Yoruba", "Hausa", "Igbo", "Zulu", "Afrikaans", "Arabic",
             "Turkish", "Russian", "Ukrainian", "Polish", "Hindi", "Bengali", "Urdu", "Chinese (Simplified)",
             "Chinese (Traditional)", "Japanese", "Korean", "Vietnamese", "Thai", "Indonesian", "Malay", "Filipino",
             "Greek", "Swedish", "Norwegian", "Danish", "Finnish", "Czech", "Romanian", "Hungarian", "Hebrew",
             "Persian"]


def ai_ready() -> bool:
    mode = paths.settings.get("ai_mode", "Auto")
    on = llm.online_config() is not None
    if mode == "Online":
        return on
    if mode == "Offline":
        return llm.LOCAL.available()
    return on or llm.LOCAL.available()


def _chunk_limit() -> int:
    """Characters of document per AI request (local models have small context windows)."""
    mode = paths.settings.get("ai_mode", "Auto")
    online = llm.online_config() is not None and mode != "Offline"
    return 24000 if online else 6000


def pdf_text(src: Path, pages: str = "all") -> list[str]:
    doc = pdfium.PdfDocument(str(src))
    out = []
    for i in parse_pages(pages, len(doc)):
        tp = doc[i].get_textpage()
        out.append(tp.get_text_range().replace("\r\n", "\n").strip())
        tp.close()
    doc.close()
    return out


def _chunks(text: str, limit: int) -> list[str]:
    paras = re.split(r"\n\s*\n|\n", text)
    out, cur = [], ""
    for p in paras:
        if len(cur) + len(p) + 1 > limit and cur:
            out.append(cur)
            cur = ""
        while len(p) > limit:
            out.append(p[:limit])
            p = p[limit:]
        cur += p + "\n"
    if cur.strip():
        out.append(cur)
    return out


# ===================================================================== extractive fallback
_STOP = set("""a an the and or but if of to in on at by for with from as is are was were be been being it its this that
these those he she they we you i his her their our your not no so than then there here which who whom what when where
how all any both each few more most other some such only own same too very can will just should now also into over
under about after before between during out up down off again further once may might must shall would could""".split())


def extractive_summary(text: str, sentences: int = 8) -> list[str]:
    sents = [s.strip() for s in re.split(r"(?<=[.!?])\s+(?=[A-Z0-9\"'(])", re.sub(r"\s+", " ", text)) if
             25 <= len(s.strip()) <= 400]
    if not sents:
        return []
    words = [w for w in re.findall(r"[a-zA-ZÀ-ɏ]{3,}", text.lower()) if w not in _STOP]
    freq = Counter(words)
    top = freq.most_common(1)[0][1] if freq else 1
    seen, scored = set(), []
    for i, s in enumerate(sents):
        key = s.lower()[:60]
        if key in seen:
            continue
        seen.add(key)
        ws = [w for w in re.findall(r"[a-zA-ZÀ-ɏ]{3,}", s.lower()) if w not in _STOP]
        if not ws:
            continue
        score = sum(freq[w] / top for w in ws) / (len(ws) ** 0.6)
        score *= 1.25 if i < len(sents) * 0.15 else 1.0          # opening sentences matter more
        scored.append((score, i, s))
    best = sorted(scored, reverse=True)[:sentences]
    return [s for _, _, s in sorted(best, key=lambda t: t[1])]


# ===================================================================== summarizer
LENGTHS = {"Short": ("5 bullet points", 5, 350), "Medium": ("8-12 bullet points", 10, 700),
           "Detailed": ("a detailed structured summary with headings and bullet points", 18, 1400)}


def summarize_text(text: str, length: str = "Medium", style: str = "Bullet points", language: str = "Same as document",
                   focus: str = "", progress=None) -> tuple[str, bool]:
    """Returns (markdown_summary, used_ai)."""
    progress = progress or (lambda f, m: None)
    shape, n_sent, max_tok = LENGTHS.get(length, LENGTHS["Medium"])
    if not text.strip():
        raise ValueError("This PDF has no selectable text. Run OCR PDF first, then summarize.")
    if not ai_ready():
        sents = extractive_summary(text, n_sent)
        return "## Key points\n\n" + "\n".join(f"- {s}" for s in sents) + "\n", False
    lang = "in the same language as the document" if language.startswith("Same") else f"in {language}"
    how = shape if style == "Bullet points" else "a few clear paragraphs"
    extra = f" Focus on: {focus}." if focus.strip() else ""
    system = ("You are an expert analyst who writes accurate, concise summaries of documents. Never invent facts. "
              "Use Markdown.")
    parts = _chunks(text, _chunk_limit())
    if len(parts) > 1:                                 # map-reduce for long documents
        notes = []
        for k, part in enumerate(parts):
            progress(0.1 + 0.6 * k / len(parts), f"Reading part {k + 1} of {len(parts)}")
            notes.append(llm.chat([{"role": "system", "content": system},
                                   {"role": "user", "content": "Write compact notes (key facts, numbers, names, "
                                    f"conclusions) of this part of a document.{extra}\n\n---\n{part}"}],
                                  max_tokens=600, temperature=0.2))
        text = "\n\n".join(notes)
    progress(0.8, "Writing summary")
    out = llm.chat([{"role": "system", "content": system},
                    {"role": "user", "content": f"Summarize the document below as {how}, {lang}. Start with a one-"
                     f"sentence overview in bold, then the summary.{extra}\n\n---\n{text[:_chunk_limit()]}"}],
                   max_tokens=max_tok, temperature=0.3)
    return out.strip() + "\n", True


def markdown_to_pdf(md: str, dest: Path, title: str = "") -> Path:
    """Small Markdown renderer (headings, bullets, bold) -> nicely set PDF."""
    from xml.sax.saxutils import escape

    from reportlab.lib.colors import HexColor
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

    font = unicode_font(script_of(md))
    bold = unicode_font(script_of(md), bold=True)
    body = ParagraphStyle("b", fontName=font, fontSize=10.5, leading=15, spaceAfter=6)
    h1 = ParagraphStyle("h1", parent=body, fontName=bold, fontSize=18, leading=23, spaceAfter=10,
                        textColor=HexColor("#1f1f3a"))
    h2 = ParagraphStyle("h2", parent=body, fontName=bold, fontSize=13, leading=18, spaceBefore=8)

    def inline(s: str) -> str:
        s = escape(s)
        s = re.sub(r"\*\*(.+?)\*\*", lambda m: f'<font name="{bold}">{m.group(1)}</font>', s)
        s = re.sub(r"(?<!\*)\*(?!\s)(.+?)(?<!\s)\*(?!\*)", r"<i>\1</i>", s)
        return s

    story = [Paragraph(inline(title), h1)] if title else []
    bullets: list = []

    def flush():
        if bullets:
            story.append(ListFlowable([ListItem(Paragraph(inline(b), body), leftIndent=12) for b in bullets],
                                      bulletType="bullet", start="•", leftIndent=14))
            bullets.clear()

    for line in md.splitlines():
        s = line.strip()
        if not s:
            flush()
            continue
        m = re.match(r"^(#{1,6})\s+(.*)", s)
        if m:
            flush()
            story.append(Paragraph(inline(m.group(2)), h1 if len(m.group(1)) == 1 and not title else h2))
        elif re.match(r"^([-*•]|\d+[.)])\s+", s):
            bullets.append(re.sub(r"^([-*•]|\d+[.)])\s+", "", s))
        else:
            flush()
            story.append(Paragraph(inline(s), body))
    flush()
    story.append(Spacer(1, 12))
    SimpleDocTemplate(str(dest), pagesize=A4, leftMargin=56, rightMargin=56, topMargin=56, bottomMargin=56,
                      title=title).build(story)
    return dest


@tool(id="ai_summarizer", name="AI Summarizer", category=INTEL, icon="ai", color="#7C5CFF", accepts=PDF, new=True,
      desc="Summarize long PDFs into concise, easy-to-read key points. Works with the online or offline AI.",
      params=[Param("length", "Summary length", "choice", "Medium", ["Short", "Medium", "Detailed"]),
              Param("style", "Style", "choice", "Bullet points", ["Bullet points", "Paragraphs"]),
              Param("language", "Language", "choice", "Same as document", ["Same as document"] + LANGUAGES),
              Param("focus", "Focus on (optional)", "text", "", hint="e.g. 'dates and money', 'risks', 'action items'"),
              Param("pages", "Pages", "pages", "all"),
              Param("output", "Save as", "choice", "PDF", ["PDF", "Word (.docx)", "Markdown (.md)", "Text (.txt)"])])
def ai_summarizer(files, out_dir, p, progress):
    src = files[0]
    progress(0.05, "Reading PDF")
    text = "\n\n".join(pdf_text(src, p.get("pages")))
    md, used_ai = summarize_text(text, p.get("length", "Medium"), p.get("style", "Bullet points"),
                                 p.get("language", "Same as document"), p.get("focus", ""), progress)
    if not used_ai:
        md = ("_Offline quick summary (key sentences). Add a free AI key or the offline AI model in Settings "
              "for a written summary._\n\n") + md
    title = f"Summary - {src.stem}"
    fmt = p.get("output", "PDF")
    if fmt.startswith("Markdown"):
        dest = out_name(src, out_dir, "summary", ".md")
        dest.write_text(f"# {title}\n\n{md}", encoding="utf-8")
    elif fmt.startswith("Text"):
        dest = out_name(src, out_dir, "summary", ".txt")
        dest.write_text(f"{title}\n\n" + re.sub(r"\*\*|__|^#+\s*", "", md, flags=re.M), encoding="utf-8")
    elif fmt.startswith("Word"):
        from docx import Document
        dest = out_name(src, out_dir, "summary", ".docx")
        d = Document()
        d.add_heading(title, 0)
        for line in md.splitlines():
            s = line.strip()
            if not s:
                continue
            m = re.match(r"^(#{1,6})\s+(.*)", s)
            if m:
                d.add_heading(m.group(2).replace("**", ""), min(3, len(m.group(1))))
            elif re.match(r"^([-*•]|\d+[.)])\s+", s):
                d.add_paragraph(re.sub(r"^([-*•]|\d+[.)])\s+", "", s).replace("**", ""), style="List Bullet")
            else:
                d.add_paragraph(s.replace("**", ""))
        d.save(str(dest))
    else:
        dest = markdown_to_pdf(md, out_name(src, out_dir, "summary"), title)
    return [dest]


# ===================================================================== ask your PDF
def ask(src: Path, question: str, history: list[dict] | None = None, on_token=None) -> str:
    """Answer a question about a PDF, citing page numbers."""
    if not ai_ready():
        raise llm.LLMError("Set up the AI first (Settings > AI): a free Gemini/Groq key or the offline model.")
    pages = pdf_text(src)
    limit = _chunk_limit()
    full = "\n\n".join(f"[Page {i + 1}]\n{t}" for i, t in enumerate(pages))
    if len(full) > limit:                       # keep the most relevant pages (keyword overlap)
        q = set(w for w in re.findall(r"\w{3,}", question.lower()) if w not in _STOP)
        ranked = sorted(range(len(pages)), key=lambda i: -sum(pages[i].lower().count(w) for w in q))
        keep, size = [], 0
        for i in ranked:
            if size + len(pages[i]) > limit:
                continue
            keep.append(i)
            size += len(pages[i])
        full = "\n\n".join(f"[Page {i + 1}]\n{pages[i]}" for i in sorted(keep))
    msgs = [{"role": "system", "content": "Answer questions using only the document below. Cite pages like (p. 3). "
             "If the answer isn't in the document, say so.\n\n" + full}]
    msgs += (history or [])[-6:]
    msgs.append({"role": "user", "content": question})
    return llm.chat(msgs, max_tokens=700, temperature=0.2, on_token=on_token)


# ===================================================================== translate
def _parse_list(text: str, n: int) -> list[str] | None:
    m = re.search(r"\[.*\]", text, re.S)
    if not m:
        return None
    try:
        arr = json.loads(m.group(0))
    except json.JSONDecodeError:
        return None
    if isinstance(arr, list) and len(arr) == n:
        return [str(x) for x in arr]
    return None


def translate_texts(texts: list[str], language: str, progress=None) -> list[str]:
    progress = progress or (lambda f, m: None)
    out: list[str] = [""] * len(texts)
    batches, cur, size = [], [], 0
    limit = min(_chunk_limit() // 2, 6000)
    for i, t in enumerate(texts):
        if cur and size + len(t) > limit:
            batches.append(cur)
            cur, size = [], 0
        cur.append(i)
        size += len(t)
    if cur:
        batches.append(cur)
    system = (f"You are a professional translator. Translate into {language}. Keep numbers, names, e-mail addresses, "
              "links and formatting. Reply ONLY with a JSON array of strings, same length and order as the input.")
    for k, b in enumerate(batches):
        progress(k / len(batches), f"Translating ({k + 1}/{len(batches)})")
        items = [texts[i] for i in b]
        res = _parse_list(llm.chat([{"role": "system", "content": system},
                                    {"role": "user", "content": json.dumps(items, ensure_ascii=False)}],
                                   max_tokens=min(8000, 400 + int(sum(len(x) for x in items) * 1.2)),
                                   temperature=0.1), len(items))
        if res is None:                            # fall back to one item at a time
            res = []
            for t in items:
                r = llm.chat([{"role": "system", "content": f"Translate into {language}. Reply with the translation "
                               "only."}, {"role": "user", "content": t}], max_tokens=1200, temperature=0.1)
                res.append(r.strip())
        for i, r in zip(b, res):
            out[i] = r
    return out


def _fit(c, text: str, font: str, box_w: float, box_h: float, start: float):
    from reportlab.lib.utils import simpleSplit
    size = start
    while size > 4:
        lines = simpleSplit(text, font, size, box_w)
        if len(lines) * size * 1.18 <= box_h * 1.12:
            return size, lines
        size -= 0.5
    return 4, simpleSplit(text, font, 4, box_w)


@tool(id="translate_pdf", name="Translate PDF", category=INTEL, icon="translate", color="#7C5CFF", accepts=PDF,
      new=True, desc="Easily translate PDF files while keeping the original layout (online or offline AI).",
      params=[Param("language", "Translate to", "choice", "English", LANGUAGES),
              Param("pages", "Pages", "pages", "all"),
              Param("output", "Output", "choice", "Translated PDF (same layout)",
                    ["Translated PDF (same layout)", "Text document (.docx)"]),
              Param("color", "Text color", "color", "#111111", show_if=("output", "Translated PDF (same layout)"))])
def translate_pdf(files, out_dir, p, progress):
    if not ai_ready():
        raise llm.LLMError("Translation needs the AI: add a free Gemini/Groq key in Settings, or install the "
                           "offline AI model.")
    src = files[0]
    lang = p.get("language", "English")
    doc = pdfium.PdfDocument(str(src))
    idx = parse_pages(p.get("pages"), len(doc))
    page_blocks = {i: textlayer.blocks(doc[i]) for i in idx}
    flat = [(i, k) for i in idx for k in range(len(page_blocks[i]))]
    if not flat:
        doc.close()
        raise ValueError("No text found to translate. If it's a scan, run OCR PDF first.")
    texts = [page_blocks[i][k]["text"] for i, k in flat]
    trans = translate_texts(texts, lang, lambda f, m: progress(0.05 + 0.75 * f, m))
    tag = re.sub(r"\W+", "", lang.split()[0])[:12].lower()
    if p.get("output", "").startswith("Text"):
        from docx import Document
        d = Document()
        cur = None
        for (i, _k), t in zip(flat, trans):
            if cur is not None and i != cur:
                d.add_page_break()
            cur = i
            d.add_paragraph(t)
        doc.close()
        dest = out_name(src, out_dir, tag, ".docx")
        d.save(str(dest))
        return [dest]

    progress(0.82, "Rebuilding pages")
    from .convert import _strip_text
    for i in idx:
        _strip_text(doc[i])
    tmp = Path(tempfile.mkdtemp()) / "stripped.pdf"
    doc.save(str(tmp))
    doc.close()
    pdf = open_pdf(tmp)
    color = hex_to_rgb(p.get("color", "#111111"))
    font = unicode_font(script_of(" ".join(trans[:20])))
    by_page: dict[int, list] = {}
    for (i, k), t in zip(flat, trans):
        b = page_blocks[i][k]
        by_page.setdefault(i, []).append((rect_to_visual(pdf.pages[i], b["rect"]), b, t))

    def draw(c, i, vw, vh):
        c.setFillColorRGB(*color)
        for (l, bt, r, tp), b, t in by_page.get(i, []):
            w = max(r - l, 40)
            if b["lines"] == 1:                      # translations run longer: allow growing to the margin
                w = max(w, min(vw - l - 24, w * 1.8))
            h = max(tp - bt, b["size"])
            size, lines = _fit(c, t, font, w, h, max(5.0, b["size"] * 0.82))
            y = tp - size
            for ln in lines:
                c.setFont(font, size)
                c.drawString(l, y, ln)
                y -= size * 1.18

    overlay_pages(pdf, draw, sorted(by_page))
    dest = out_name(src, out_dir, tag)
    save(pdf, dest)
    return [dest]
