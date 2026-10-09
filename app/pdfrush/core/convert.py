"""Conversions to and from PDF."""
from __future__ import annotations

import io
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

import pikepdf
import pypdfium2 as pdfium

from .. import paths
from . import textlayer
from .registry import IMAGES, OFFICE_PPT, OFFICE_WORD, OFFICE_XLS, PDF, Param, tool
from .util import open_pdf, out_name, parse_pages, unique

TO = "Convert to PDF"
FROM = "Convert from PDF"
_NO_WINDOW = 0x08000000 if sys.platform == "win32" else 0
PAGE_SIZES = {"A4": (595.28, 841.89), "Letter": (612, 792), "Legal": (612, 1008), "A3": (841.89, 1190.55),
              "A5": (419.53, 595.28)}


def _pdfium_doc(path: Path) -> pdfium.PdfDocument:
    try:
        return pdfium.PdfDocument(str(path))
    except pdfium.PdfiumError as e:
        if "password" in str(e).lower():
            raise ValueError(f"'{path.name}' is password-protected - unlock it first.") from None
        raise


# ===================================================================== PDF -> images
@tool(id="pdf_to_jpg", name="PDF to JPG", category=FROM, icon="image", color="#FFC857", accepts=PDF,
      desc="Convert each PDF page into a JPG or extract all images contained in a PDF.",
      params=[Param("mode", "Convert", "choice", "Pages to images", ["Pages to images", "Extract images"]),
              Param("format", "Format", "choice", "JPG", ["JPG", "PNG"], show_if=("mode", "Pages to images")),
              Param("dpi", "Quality (DPI)", "choice", "150", ["72", "150", "220", "300", "600"],
                    show_if=("mode", "Pages to images")),
              Param("pages", "Pages", "pages", "all", show_if=("mode", "Pages to images")),
              Param("zip", "Put everything in one ZIP", "bool", False)])
def pdf_to_jpg(files, out_dir, p, progress):
    src = files[0]
    folder = unique(Path(out_dir) / f"{src.stem}_images")
    folder.mkdir(parents=True)
    outs = []
    if p["mode"] == "Extract images":
        pdf = open_pdf(src)
        seen = set()
        for pi, page in enumerate(pdf.pages):
            progress(pi / max(1, len(pdf.pages)), f"Page {pi + 1}")
            for name, raw in page.images.items():
                if raw.objgen in seen:
                    continue
                seen.add(raw.objgen)
                try:
                    img = pikepdf.PdfImage(raw)
                    target = folder / f"page{pi + 1}_image{len(seen)}"
                    outs.append(Path(img.extract_to(fileprefix=str(target))))
                except Exception:
                    continue
        if not outs:
            raise ValueError("No embedded images found in this PDF.")
    else:
        doc = _pdfium_doc(src)
        scale = int(p.get("dpi", 150)) / 72
        ext = "jpg" if p.get("format", "JPG") == "JPG" else "png"
        idx = parse_pages(p.get("pages"), len(doc))
        width = len(str(len(doc)))
        for k, i in enumerate(idx):
            progress(k / max(1, len(idx)), f"Rendering page {i + 1}")
            img = doc[i].render(scale=scale, may_draw_forms=True).to_pil()
            dest = folder / f"{src.stem}_page{str(i + 1).zfill(width)}.{ext}"
            if ext == "jpg":
                img.convert("RGB").save(dest, "JPEG", quality=90, optimize=True)
            else:
                img.save(dest, "PNG", optimize=True)
            outs.append(dest)
        doc.close()
    if p.get("zip"):
        z = unique(Path(out_dir) / f"{src.stem}_images.zip")
        with zipfile.ZipFile(z, "w", zipfile.ZIP_DEFLATED) as zf:
            for f in outs:
                zf.write(f, f.name)
        shutil.rmtree(folder, ignore_errors=True)
        return [z]
    return outs


# ===================================================================== images -> PDF
def _image_pages(files, page="Fit to image", orientation="Auto", margin=0, quality=90):
    from PIL import Image, ImageOps
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    for f in files:
        im = Image.open(f)
        frames = []
        try:
            for k in range(getattr(im, "n_frames", 1)):
                im.seek(k)
                frames.append(ImageOps.exif_transpose(im.copy()))
        except EOFError:
            pass
        for fr in frames or [im]:
            if fr.mode not in ("RGB", "L", "1"):
                fr = fr.convert("RGB")
            iw, ih = fr.size
            dpi = fr.info.get("dpi", (96, 96))[0] or 96
            if page == "Fit to image":
                pw, ph = iw * 72 / dpi + 2 * margin, ih * 72 / dpi + 2 * margin
            else:
                pw, ph = PAGE_SIZES[page]
                land = (orientation == "Landscape") or (orientation == "Auto" and iw > ih)
                if land:
                    pw, ph = ph, pw
            c.setPageSize((pw, ph))
            aw, ah = pw - 2 * margin, ph - 2 * margin
            s = min(aw / iw, ah / ih)
            w, h = iw * s, ih * s
            b = io.BytesIO()
            fr.save(b, "JPEG" if fr.mode == "RGB" else "PNG", quality=quality)
            b.seek(0)
            c.drawImage(ImageReader(b), (pw - w) / 2, (ph - h) / 2, w, h)
            c.showPage()
    c.save()
    buf.seek(0)
    return buf


def images_to_pdf_bytes(files) -> io.BytesIO:
    return _image_pages(files)


@tool(id="jpg_to_pdf", name="JPG to PDF", category=TO, icon="image", color="#FFC857", accepts=IMAGES, combine=True,
      desc="Convert JPG images to PDF in seconds. Easily adjust orientation and margins.",
      params=[Param("page", "Page size", "choice", "A4", ["A4", "Letter", "Legal", "A3", "A5", "Fit to image"]),
              Param("orientation", "Orientation", "choice", "Auto", ["Auto", "Portrait", "Landscape"]),
              Param("margin", "Margin", "choice", "Small", ["None", "Small", "Big"]),
              Param("one_file", "Merge all images into one PDF", "bool", True),
              Param("name", "Output name", "text", "images")])
def jpg_to_pdf(files, out_dir, p, progress):
    margin = {"None": 0, "Small": 20, "Big": 50}[p.get("margin", "Small")]
    groups = [files] if p.get("one_file", True) else [[f] for f in files]
    outs = []
    for k, g in enumerate(groups):
        progress(k / len(groups), f"Building PDF {k + 1}")
        buf = _image_pages(g, p.get("page", "A4"), p.get("orientation", "Auto"), margin)
        dest = unique(Path(out_dir) / (f"{p.get('name') or 'images'}.pdf" if len(groups) == 1 else f"{g[0].stem}.pdf"))
        dest.write_bytes(buf.getvalue())
        outs.append(dest)
    return outs


# ===================================================================== PDF -> Word
@tool(id="pdf_to_word", name="PDF to Word", category=FROM, icon="word", color="#2B6CB0", accepts=PDF,
      desc="Easily convert your PDF files into easy to edit DOC and DOCX documents.",
      params=[Param("pages", "Pages", "pages", "all"),
              Param("ocr", "Scanned PDF? Recognise text first (OCR)", "bool", False)])
def pdf_to_word(files, out_dir, p, progress):
    src = files[0]
    if p.get("ocr"):
        from .ocr import ocr_pdf_file
        tmp = Path(tempfile.mkdtemp()) / src.name
        ocr_pdf_file(src, tmp, progress=lambda f, m: progress(0.4 * f, m))
        src = tmp
    dest = out_name(files[0], out_dir, "", ".docx")
    try:
        from pdf2docx import Converter     # best layout fidelity (tables, columns, images)
        doc = _pdfium_doc(src)
        n = len(doc)
        doc.close()
        idx = parse_pages(p.get("pages"), n)
        progress(0.45, "Converting layout")
        cv = Converter(str(src))
        cv.convert(str(dest), pages=idx)
        cv.close()
        return [dest]
    except ImportError:
        pass
    return [_simple_docx(src, dest, p.get("pages"), progress)]


def _simple_docx(src: Path, dest: Path, pages: str, progress) -> Path:
    from docx import Document
    from docx.shared import Pt
    doc = _pdfium_doc(src)
    out = Document()
    idx = parse_pages(pages, len(doc))
    for k, i in enumerate(idx):
        progress(0.5 + 0.5 * k / max(1, len(idx)), f"Page {i + 1}")
        bl = textlayer.blocks(doc[i])
        sizes = sorted(b["size"] for b in bl) or [10]
        body = sizes[len(sizes) // 2]
        for b in bl:
            if b["lines"] == 1 and b["size"] > body * 1.35:
                out.add_heading(b["text"], level=1 if b["size"] > body * 1.8 else 2)
            else:
                para = out.add_paragraph(b["text"])
                for r in para.runs:
                    r.font.size = Pt(max(8, min(14, round(b["size"] * 0.8))))
        if k < len(idx) - 1:
            out.add_page_break()
    doc.close()
    out.save(str(dest))
    return dest


def _strip_text(page) -> None:
    """Remove all top-level text objects from a pdfium page (keeps images/vector art)."""
    gone = [o for o in page.get_objects(max_depth=1) if o.type == pdfium.raw.FPDF_PAGEOBJ_TEXT]
    for o in gone:
        page.remove_obj(o)
    page.gen_content()
    for o in gone:
        try:
            o.close()
        except Exception:
            pass


# ===================================================================== PDF -> PowerPoint
@tool(id="pdf_to_ppt", name="PDF to PowerPoint", category=FROM, icon="ppt", color="#DD6B20", accepts=PDF,
      desc="Turn your PDF files into easy to edit PPT and PPTX slideshows.",
      params=[Param("mode", "Slides", "choice", "Editable text", ["Editable text", "Exact look (images)"]),
              Param("pages", "Pages", "pages", "all")])
def pdf_to_ppt(files, out_dir, p, progress):
    from pptx import Presentation
    from pptx.util import Emu, Pt
    src = files[0]
    doc = _pdfium_doc(src)
    idx = parse_pages(p.get("pages"), len(doc))
    w0, h0 = doc[idx[0]].get_size() if idx else (960, 540)
    if doc[idx[0]].get_rotation() in (90, 270):
        w0, h0 = h0, w0
    prs = Presentation()
    EMU = 12700
    prs.slide_width, prs.slide_height = Emu(int(w0 * EMU)), Emu(int(h0 * EMU))
    editable = p.get("mode", "").startswith("Editable")
    tmpdir = Path(tempfile.mkdtemp())
    for k, i in enumerate(idx):
        progress(k / max(1, len(idx)), f"Slide {k + 1}")
        page = doc[i]
        W, H = page.get_size()
        rot = page.get_rotation()
        blocks = textlayer.blocks(page) if editable and rot == 0 else []
        if blocks:
            _strip_text(page)
        img = page.render(scale=2).to_pil().convert("RGB")
        page.close()
        bg = tmpdir / f"s{k}.jpg"
        img.save(bg, "JPEG", quality=88)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        sw, sh = prs.slide_width, prs.slide_height
        slide.shapes.add_picture(str(bg), 0, 0, sw, sh)
        fx, fy = sw / (W if rot in (0, 180) else H), sh / (H if rot in (0, 180) else W)
        for b in blocks:
            x0, y0, x1, y1 = b["rect"]
            tb = slide.shapes.add_textbox(Emu(int(x0 * fx)), Emu(int((H - y1) * fy)), Emu(int((x1 - x0) * fx) + 20000),
                                          Emu(int((y1 - y0) * fy) + 20000))
            tf = tb.text_frame
            tf.word_wrap = False
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            tf.text = "\n".join(b["rows"])
            for para in tf.paragraphs:
                for r in para.runs:
                    r.font.size = Pt(max(6, round(b["size"] * 0.82)))
    doc.close()
    dest = out_name(src, out_dir, "", ".pptx")
    prs.save(str(dest))
    shutil.rmtree(tmpdir, ignore_errors=True)
    return [dest]


# ===================================================================== PDF -> Excel
@tool(id="pdf_to_excel", name="PDF to Excel", category=FROM, icon="excel", color="#2F855A", accepts=PDF,
      desc="Pull data straight from PDFs into Excel spreadsheets in a few short seconds.",
      params=[Param("layout", "Sheets", "choice", "One sheet per table", ["One sheet per table", "One sheet per page",
                                                                         "Everything in one sheet"]),
              Param("pages", "Pages", "pages", "all")])
def pdf_to_excel(files, out_dir, p, progress):
    import pdfplumber
    from openpyxl import Workbook
    from openpyxl.styles import Font
    src = files[0]
    wb = Workbook()
    wb.remove(wb.active)
    layout = p.get("layout", "One sheet per table")
    found = 0
    one = wb.create_sheet("Data") if layout.startswith("Everything") else None
    with pdfplumber.open(str(src)) as pdf:
        idx = parse_pages(p.get("pages"), len(pdf.pages))
        for k, i in enumerate(idx):
            progress(k / max(1, len(idx)), f"Page {i + 1}")
            page = pdf.pages[i]
            tables = page.extract_tables() or []
            if not tables:                       # no ruled table: split text lines on wide gaps
                rows = [re.split(r"\s{2,}", ln.strip()) for ln in (page.extract_text(layout=True) or "").splitlines()
                        if ln.strip()]
                if rows and max(len(r) for r in rows) > 1:
                    tables = [rows]
            ws_page = wb.create_sheet(f"Page {i + 1}") if layout.endswith("per page") and tables else None
            for t_i, table in enumerate(tables):
                found += 1
                ws = one or ws_page or wb.create_sheet(f"P{i + 1} T{t_i + 1}")
                if ws.max_row > 1 or (ws.max_row == 1 and ws.cell(1, 1).value is not None):
                    ws.append([])
                for r_i, row in enumerate(table):
                    vals = [_num(c) for c in row]
                    ws.append(vals)
                    if r_i == 0:
                        for cell in ws[ws.max_row]:
                            cell.font = Font(bold=True)
    if not found:
        raise ValueError("No tables found. If the PDF is a scan, run OCR PDF first.")
    for ws in wb.worksheets:
        for col in ws.columns:
            width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[col[0].column_letter].width = min(60, max(8, width + 2))
    dest = out_name(src, out_dir, "", ".xlsx")
    wb.save(str(dest))
    return [dest]


def _num(v):
    if v is None:
        return None
    s = str(v).replace("\n", " ").strip()
    t = s.replace(",", "").replace(" ", "")
    if re.fullmatch(r"-?\d+", t) and len(t) < 15 and not (t.startswith("0") and len(t) > 1):
        return int(t)
    if re.fullmatch(r"-?\d*\.\d+", t):
        try:
            return float(t)
        except ValueError:
            pass
    return s


# ===================================================================== Office -> PDF
def _soffice() -> str | None:
    cands = [shutil.which("soffice"), shutil.which("libreoffice")]
    for base in (os.environ.get("ProgramFiles", r"C:\Program Files"), os.environ.get("ProgramFiles(x86)", "")):
        if base:
            cands.append(str(Path(base) / "LibreOffice" / "program" / "soffice.exe"))
    cands.append(str(paths.ROOT_DIR / "libreoffice" / "program" / "soffice.exe"))
    return next((c for c in cands if c and Path(c).exists()), None)


def _ms_office(src: Path, dest: Path) -> bool:
    if sys.platform != "win32":
        return False
    try:
        import comtypes.client as cc
    except ImportError:
        return False
    ext = src.suffix.lower()
    s, d = str(src.resolve()), str(dest.resolve())
    try:
        if ext in (".doc", ".docx", ".odt", ".rtf", ".txt"):
            app = cc.CreateObject("Word.Application")
            app.Visible = False
            try:
                doc = app.Documents.Open(s, ReadOnly=True)
                doc.SaveAs(d, FileFormat=17)
                doc.Close(False)
            finally:
                app.Quit()
        elif ext in (".ppt", ".pptx", ".odp"):
            app = cc.CreateObject("PowerPoint.Application")
            try:
                pres = app.Presentations.Open(s, ReadOnly=True, WithWindow=False)
                pres.SaveAs(d, 32)
                pres.Close()
            finally:
                app.Quit()
        elif ext in (".xls", ".xlsx", ".ods", ".csv"):
            app = cc.CreateObject("Excel.Application")
            app.Visible = False
            app.DisplayAlerts = False
            try:
                wb = app.Workbooks.Open(s, ReadOnly=True)
                wb.ExportAsFixedFormat(0, d)
                wb.Close(False)
            finally:
                app.Quit()
        else:
            return False
        return dest.exists()
    except Exception:
        return False


def office_to_pdf(src: Path, out_dir: Path) -> Path:
    dest = out_name(src, out_dir, "")
    engine = paths.settings.get("office_engine", "Auto")
    if engine in ("Auto", "Microsoft Office") and _ms_office(src, dest):
        return dest
    so = _soffice()
    if engine in ("Auto", "LibreOffice") and so:
        tmp = Path(tempfile.mkdtemp())
        r = subprocess.run([so, "--headless", "--norestore", "--convert-to", "pdf", "--outdir", str(tmp), str(src)],
                           capture_output=True, timeout=600, creationflags=_NO_WINDOW)
        made = tmp / (src.stem + ".pdf")
        if made.exists():
            shutil.move(str(made), dest)
            return dest
        raise RuntimeError(f"LibreOffice couldn't convert {src.name}: {r.stderr.decode(errors='ignore')[:300]}")
    if src.suffix.lower() in (".txt",):
        return _text_to_pdf(src, dest)
    raise RuntimeError("To convert Office files, PDFRush uses Microsoft Office (if installed) or the free "
                       "LibreOffice. Install LibreOffice from libreoffice.org, then try again.")


def _text_to_pdf(src: Path, dest: Path) -> Path:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate
    from xml.sax.saxutils import escape
    text = src.read_text(encoding="utf-8", errors="replace")
    st = getSampleStyleSheet()["BodyText"]
    story = [Paragraph(escape(p).replace("\n", "<br/>"), st) for p in text.split("\n\n")]
    SimpleDocTemplate(str(dest), pagesize=A4).build(story)
    return dest


for _id, _name, _acc, _icon, _color, _desc in (
        ("word_to_pdf", "Word to PDF", OFFICE_WORD, "word", "#2B6CB0",
         "Make DOC and DOCX files easy to read by converting them to PDF."),
        ("ppt_to_pdf", "PowerPoint to PDF", OFFICE_PPT, "ppt", "#DD6B20",
         "Make PPT and PPTX slideshows easy to view by converting them to PDF."),
        ("excel_to_pdf", "Excel to PDF", OFFICE_XLS, "excel", "#2F855A",
         "Make EXCEL spreadsheets easy to read by converting them to PDF.")):
    tool(id=_id, name=_name, category=TO, icon=_icon, color=_color, accepts=_acc, desc=_desc)(
        lambda files, out_dir, p, progress: [office_to_pdf(files[0], out_dir)])


# ===================================================================== HTML -> PDF
def _browser() -> str | None:
    cands = []
    for base in (os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)"),
                 os.environ.get("ProgramFiles", r"C:\Program Files"), os.environ.get("LOCALAPPDATA", "")):
        if base:
            cands += [Path(base) / "Microsoft/Edge/Application/msedge.exe",
                      Path(base) / "Google/Chrome/Application/chrome.exe"]
    for n in ("msedge", "google-chrome", "chromium", "chromium-browser", "chrome"):
        w = shutil.which(n)
        if w:
            cands.append(Path(w))
    pw = Path(os.environ.get("PLAYWRIGHT_BROWSERS_PATH", "/opt/pw-browsers"))
    if pw.exists():
        cands += list(pw.glob("chromium*/chrome-linux*/chrome"))
    return next((str(c) for c in cands if Path(c).exists()), None)


@tool(id="html_to_pdf", name="HTML to PDF", category=TO, icon="web", color="#FFC857", accepts=[".html", ".htm", ".mht"],
      url_input=True,
      desc="Convert webpages in HTML to PDF. Copy and paste the URL of the page you want and convert it to PDF "
           "with a click.",
      params=[Param("url", "Web address (or drop .html files)", "text", "https://"),
              Param("background", "Include background colors/images", "bool", True)])
def html_to_pdf(files, out_dir, p, progress):
    br = _browser()
    if not br:
        raise RuntimeError("Microsoft Edge or Google Chrome is needed to convert web pages.")
    targets = [(f.resolve().as_uri(), f.stem) for f in files]
    url = str(p.get("url", "")).strip()
    if url and url not in ("https://", "http://"):
        url = url if re.match(r"^[a-z]+://", url) else "https://" + url
        targets.append((url, re.sub(r"[^\w.-]+", "_", re.sub(r"^\w+://", "", url))[:60].strip("_")))
    if not targets:
        raise ValueError("Paste a web address or add .html files.")
    outs = []
    for k, (t, name) in enumerate(targets):
        progress(k / len(targets), f"Loading {t[:60]}")
        dest = unique(Path(out_dir) / f"{name or 'page'}.pdf")
        user = tempfile.mkdtemp()
        cmd = [br, "--headless=new", "--disable-gpu", "--no-sandbox", f"--user-data-dir={user}",
               "--no-pdf-header-footer", "--print-to-pdf-no-header", "--run-all-compositor-stages-before-draw",
               "--virtual-time-budget=8000", f"--print-to-pdf={dest}", t]
        if not p.get("background", True):
            cmd.insert(1, "--disable-background-graphics")
        subprocess.run(cmd, capture_output=True, timeout=180, creationflags=_NO_WINDOW)
        shutil.rmtree(user, ignore_errors=True)
        if not dest.exists():
            raise RuntimeError(f"Couldn't load {t}")
        outs.append(dest)
    return outs


# ===================================================================== PDF -> Markdown / text
@tool(id="pdf_to_markdown", name="PDF to Markdown", category=FROM, icon="markdown", color="#4A5568", accepts=PDF,
      new=True, desc="Easily turn PDFs into Markdown files. Perfect for notes, docs, and LLMs. Headings, tables, "
                     "lists, and links preserved automatically.",
      params=[Param("pages", "Pages", "pages", "all")])
def pdf_to_markdown(files, out_dir, p, progress):
    src = files[0]
    dest = out_name(src, out_dir, "", ".md")
    try:
        import pymupdf4llm
        doc = _pdfium_doc(src)
        n = len(doc)
        doc.close()
        md = pymupdf4llm.to_markdown(str(src), pages=parse_pages(p.get("pages"), n))
        dest.write_text(md, encoding="utf-8")
        return [dest]
    except ImportError:
        pass
    dest.write_text(markdown_from_pdf(src, p.get("pages"), progress), encoding="utf-8")
    return [dest]


def markdown_from_pdf(src: Path, pages="all", progress=None) -> str:
    import pdfplumber
    progress = progress or (lambda f, m: None)
    doc = _pdfium_doc(src)
    out = []
    idx = parse_pages(pages, len(doc))
    with pdfplumber.open(str(src)) as plumb:
        for k, i in enumerate(idx):
            progress(k / max(1, len(idx)), f"Page {i + 1}")
            bl = textlayer.blocks(doc[i])
            sizes = sorted(b["size"] for b in bl) or [10]
            body = sizes[len(sizes) // 2]
            tables = []
            try:
                for t in plumb.pages[i].find_tables():
                    x0, top, x1, bottom = t.bbox
                    H = float(plumb.pages[i].height)
                    tables.append(((x0, H - bottom, x1, H - top), t.extract()))
            except Exception:
                pass
            placed = set()
            for b in bl:
                hit = next((k2 for k2, (r, _) in enumerate(tables) if textlayer.intersects(b["rect"], r)), None)
                if hit is not None:
                    if hit not in placed:
                        placed.add(hit)
                        rows = [[(c or "").replace("\n", " ").replace("|", "\\|") for c in row] for row in tables[hit][1]]
                        if rows:
                            out.append("| " + " | ".join(rows[0]) + " |")
                            out.append("|" + "---|" * len(rows[0]))
                            out += ["| " + " | ".join(r) + " |" for r in rows[1:]]
                            out.append("")
                    continue
                t = b["text"].strip()
                t = re.sub(r"(https?://\S+)", r"<\1>", t)
                if b["lines"] <= 2 and b["size"] > body * 1.6:
                    out.append(f"# {t}\n")
                elif b["lines"] <= 2 and b["size"] > body * 1.25:
                    out.append(f"## {t}\n")
                elif re.match(r"^[•●▪\-–*]\s+", t):
                    out.append("- " + re.sub(r"^[•●▪\-–*]\s+", "", t))
                elif re.match(r"^\d+[.)]\s+", t):
                    out.append(t)
                else:
                    out.append(t + "\n")
            out.append("")
    doc.close()
    return "\n".join(out).strip() + "\n"


@tool(id="pdf_to_text", name="PDF to Text", category=FROM, icon="text", color="#4A5568", accepts=PDF,
      desc="Get all the text out of a PDF as a plain .txt file.",
      params=[Param("pages", "Pages", "pages", "all")])
def pdf_to_text(files, out_dir, p, progress):
    doc = _pdfium_doc(files[0])
    parts = []
    for i in parse_pages(p.get("pages"), len(doc)):
        tp = doc[i].get_textpage()
        parts.append(tp.get_text_range().replace("\r\n", "\n"))
        tp.close()
    doc.close()
    dest = out_name(files[0], out_dir, "", ".txt")
    dest.write_text("\n\n".join(parts), encoding="utf-8")
    return [dest]
